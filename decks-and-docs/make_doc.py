#!/usr/bin/env python3
"""Документ .docx (и PDF) из того же описания, что и презентация.

Зачем. Модель агента файлы сама не собирает: документ делает этот скрипт,
агент только пишет описание и отправляет готовый файл в чат.

Описание это тот же JSON, что у make_deck.py (файл или стандартный вход).
Слайды становятся разделами, поле slides можно назвать sections:

  {
    "title": "Название документа",
    "subtitle": "подзаголовок, необязательно",
    "theme": "sand",
    "sections": [
      {"title": "Раздел", "text": "абзац\\n\\nещё абзац"},
      {"title": "Раздел", "bullets": ["пункт", "ещё пункт"]},
      {"title": "Раздел", "image": "/путь/картинка.png"}
    ]
  }

Поля раздела: title, subtitle, text, bullets (список строк), image (путь).
Поле notes это заметки докладчика для презентации, в документ оно не идёт.

Запуск:
  python3 make_doc.py spec.json [--out файл.docx] [--pdf] [--theme night] [--json]
  cat spec.json | python3 make_doc.py - --pdf

PDF собирается вторым файлом рядом с .docx через reportlab, шрифт с кириллицей
берётся из системных (DejaVu или Liberation на Linux, Verdana или Arial Unicode
на маке).
Кириллического шрифта нет: PDF не делается, скрипт честно об этом пишет, .docx
всё равно готов. LibreOffice не нужен: .docx в PDF не конвертируется, PDF
собирается из описания заново.

Куда кладётся файл по умолчанию: ~/.hermes/cache/documents. Из этой папки
гейтвею разрешено отправлять вложения. Отправить в чат:
  hermes send --to telegram "MEDIA:<путь к файлу>"
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches as DocxInches
    from docx.shared import Pt as DocxPt
    from docx.shared import RGBColor as DocxColor
except ImportError:  # pragma: no cover
    sys.stderr.write(
        "Нет библиотеки python-docx. Поставить: pip install python-docx\n"
    )
    raise SystemExit(3)

# Те же три темы, что у презентаций: цвета не выдумывать.
THEMES = {
    "sand": {"ink": "1F1B16", "accent": "C0703C", "muted": "6E6259"},
    "night": {"ink": "1A1C20", "accent": "2F6FBF", "muted": "5A6470"},
    "mint": {"ink": "14231C", "accent": "2E9E6B", "muted": "5C7268"},
}
FONT = "Arial"

# Шрифты с кириллицей для PDF: первый найденный и берётся.
PDF_FONTS = [
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ("/System/Library/Fonts/Supplemental/Verdana.ttf",
     "/System/Library/Fonts/Supplemental/Verdana Bold.ttf"),
    ("/Library/Fonts/Arial Unicode.ttf", "/Library/Fonts/Arial Unicode.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
     "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
]

TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e",
    "ж": "zh", "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m",
    "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u",
    "ф": "f", "х": "h", "ц": "c", "ч": "ch", "ш": "sh", "щ": "sch", "ъ": "",
    "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
}


def slugify(text, fallback="dokument"):
    out = []
    for ch in (text or "").lower():
        out.append(TRANSLIT.get(ch, ch))
    slug = re.sub(r"[^a-z0-9]+", "-", "".join(out)).strip("-")
    return (slug or fallback)[:60]


def fail(message):
    sys.stderr.write(message.rstrip() + "\n")
    raise SystemExit(2)


def load_spec(source):
    try:
        raw = sys.stdin.read() if source == "-" else open(source, encoding="utf-8").read()
    except OSError as exc:
        fail("Не смог прочитать описание документа: %s" % exc)
    try:
        spec = json.loads(raw)
    except json.JSONDecodeError as exc:
        fail("Описание документа не разобралось как JSON: %s" % exc)
    if not isinstance(spec, dict):
        fail("Описание должно быть объектом JSON с полем sections или slides.")
    sections = spec.get("sections")
    if sections is None:
        sections = spec.get("slides")
    if not isinstance(sections, list) or not sections:
        fail("В описании нет непустого списка sections (или slides).")
    for i, item in enumerate(sections, 1):
        if not isinstance(item, dict):
            fail("Раздел %d описан не объектом JSON." % i)
        bullets = item.get("bullets")
        if bullets is not None and not isinstance(bullets, list):
            fail("Раздел %d: bullets должен быть списком строк." % i)
    theme = spec.get("theme") or "sand"
    if theme not in THEMES:
        fail("Тема «%s» неизвестна. Есть: %s" % (theme, ", ".join(sorted(THEMES))))
    spec["_sections"] = sections
    return spec


def paragraphs_of(text):
    return [p.strip() for p in re.split(r"\n\s*\n", text or "") if p.strip()]


def default_out_dir():
    home = os.environ.get("HERMES_HOME") or os.path.join(os.path.expanduser("~"), ".hermes")
    return os.path.join(home, "cache", "documents")


def style_run(run, size, color, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = DocxPt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = DocxColor.from_string(color)


def build_docx(spec, theme, out_path):
    doc = docx.Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = DocxPt(11)

    title = (spec.get("title") or "").strip()
    if title:
        par = doc.add_paragraph()
        style_run(par.add_run(title), 22, theme["ink"], bold=True)
        par.paragraph_format.space_after = DocxPt(4)
    subtitle = (spec.get("subtitle") or "").strip()
    if subtitle:
        par = doc.add_paragraph()
        style_run(par.add_run(subtitle), 12, theme["muted"], italic=True)
    if title or subtitle:
        par = doc.add_paragraph()
        style_run(par.add_run(dt.date.today().strftime("%d.%m.%Y")), 10, theme["muted"])

    for item in spec["_sections"]:
        head = (item.get("title") or "").strip()
        if head:
            par = doc.add_heading(level=1)
            style_run(par.add_run(head), 16, theme["accent"], bold=True)
        sub = (item.get("subtitle") or "").strip()
        if sub:
            par = doc.add_paragraph()
            style_run(par.add_run(sub), 11, theme["muted"], italic=True)
        for chunk in paragraphs_of(item.get("text")):
            par = doc.add_paragraph()
            style_run(par.add_run(chunk), 11, theme["ink"])
        for bullet in (item.get("bullets") or []):
            bullet = str(bullet).strip()
            if not bullet:
                continue
            par = doc.add_paragraph(style="List Bullet")
            style_run(par.add_run(bullet), 11, theme["ink"])
        image = (item.get("image") or "").strip()
        if image:
            full = os.path.expanduser(image)
            if not os.path.isfile(full):
                sys.stderr.write("ВНИМАНИЕ: картинка не найдена, раздел без неё: %s\n" % image)
            else:
                try:
                    doc.add_picture(full, width=DocxInches(6.0))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception as exc:
                    sys.stderr.write("ВНИМАНИЕ: картинку не вставил (%s): %s\n" % (exc, image))

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return os.path.abspath(out_path)


def pdf_font():
    """Первый системный шрифт с кириллицей. Нет ни одного: None."""
    for regular, bold in PDF_FONTS:
        if os.path.isfile(regular) and os.path.isfile(bold):
            return regular, bold
    return None


def build_pdf(spec, theme, out_path):
    """PDF из того же описания. Вернёт путь или None, если собрать нечем."""
    try:
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer
    except ImportError:
        sys.stderr.write("PDF не собран: нет reportlab. Поставить: pip install reportlab\n")
        return None

    found = pdf_font()
    if not found:
        sys.stderr.write("PDF не собран: в системе нет шрифта с кириллицей (искал DejaVu, Liberation, Verdana, Arial Unicode).\n")
        return None
    regular, bold = found
    pdfmetrics.registerFont(TTFont("DocSans", regular))
    pdfmetrics.registerFont(TTFont("DocSans-Bold", bold))

    from xml.sax.saxutils import escape

    def esc(text):
        return escape(str(text)).replace("\n", "<br/>")

    ink = "#" + theme["ink"]
    accent = "#" + theme["accent"]
    muted = "#" + theme["muted"]
    st_title = ParagraphStyle("t", fontName="DocSans-Bold", fontSize=20, leading=25,
                              textColor=ink, spaceAfter=6, alignment=TA_LEFT)
    st_sub = ParagraphStyle("s", fontName="DocSans", fontSize=11, leading=15,
                            textColor=muted, spaceAfter=14)
    st_head = ParagraphStyle("h", fontName="DocSans-Bold", fontSize=14, leading=18,
                             textColor=accent, spaceBefore=14, spaceAfter=6)
    st_body = ParagraphStyle("b", fontName="DocSans", fontSize=11, leading=16,
                             textColor=ink, spaceAfter=8)
    # bulletFontName обязательно тот же зарегистрированный шрифт: со шрифтом по
    # умолчанию точка списка в PDF не рисовалась вовсе.
    st_bullet = ParagraphStyle("u", parent=st_body, leftIndent=14, bulletIndent=2,
                               spaceAfter=5, bulletFontName="DocSans",
                               bulletFontSize=11)

    story = []
    title = (spec.get("title") or "").strip()
    if title:
        story.append(Paragraph(esc(title), st_title))
    subtitle = (spec.get("subtitle") or "").strip()
    if subtitle:
        story.append(Paragraph(esc(subtitle), st_sub))
    if title or subtitle:
        story.append(Paragraph(dt.date.today().strftime("%d.%m.%Y"), st_sub))

    max_w = A4[0] - 4 * cm
    for item in spec["_sections"]:
        head = (item.get("title") or "").strip()
        if head:
            story.append(Paragraph(esc(head), st_head))
        sub = (item.get("subtitle") or "").strip()
        if sub:
            story.append(Paragraph(esc(sub), st_sub))
        for chunk in paragraphs_of(item.get("text")):
            story.append(Paragraph(esc(chunk), st_body))
        for bullet in (item.get("bullets") or []):
            bullet = str(bullet).strip()
            if bullet:
                story.append(Paragraph(esc(bullet), st_bullet, bulletText="•"))
        image = (item.get("image") or "").strip()
        if image:
            full = os.path.expanduser(image)
            if not os.path.isfile(full):
                sys.stderr.write("ВНИМАНИЕ: картинка не найдена, раздел PDF без неё: %s\n" % image)
            else:
                try:
                    iw, ih = ImageReader(full).getSize()
                    width = min(max_w, iw)
                    story.append(Spacer(1, 6))
                    story.append(Image(full, width=width, height=width * ih / iw))
                    story.append(Spacer(1, 6))
                except Exception as exc:
                    sys.stderr.write("ВНИМАНИЕ: картинку в PDF не вставил (%s): %s\n" % (exc, image))

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    SimpleDocTemplate(out_path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm,
                      title=title or "Документ").build(story)
    return os.path.abspath(out_path)


def main():
    ap = argparse.ArgumentParser(description="Документ .docx (и PDF) из описания")
    ap.add_argument("spec", help="файл с описанием или - для стандартного входа")
    ap.add_argument("--out", help="куда сохранить .docx")
    ap.add_argument("--pdf", action="store_true", help="собрать ещё и PDF рядом")
    ap.add_argument("--theme", choices=sorted(THEMES), help="тема оформления")
    ap.add_argument("--json", action="store_true", help="ответ машинным JSON")
    args = ap.parse_args()

    spec = load_spec(args.spec)
    theme = THEMES[args.theme or spec.get("theme") or "sand"]

    out_path = args.out
    if not out_path:
        stamp = dt.date.today().strftime("%Y-%m-%d")
        name = "%s-%s.docx" % (slugify(spec.get("title")), stamp)
        out_path = os.path.join(default_out_dir(), name)
    out_path = os.path.expanduser(out_path)

    docx_path = build_docx(spec, theme, out_path)
    pdf_path = None
    if args.pdf:
        pdf_path = build_pdf(spec, theme, re.sub(r"\.docx$", "", docx_path) + ".pdf")

    if args.json:
        print(json.dumps({"docx": docx_path, "pdf": pdf_path,
                          "sections": len(spec["_sections"])}, ensure_ascii=False))
    else:
        print(docx_path)
        if pdf_path:
            print(pdf_path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
