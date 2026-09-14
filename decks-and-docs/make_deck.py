#!/usr/bin/env python3
"""Презентация .pptx из простого описания слайдов.

Зачем. Модель агента файлы сама не собирает: презентацию делает этот скрипт,
агент только пишет описание слайдов и отправляет готовый файл в чат.
На выходе всегда путь к файлу, а не рассказ о том, как всё устроено.

Описание слайдов это JSON (файл или стандартный вход):

  {
    "title": "Название презентации",
    "subtitle": "подзаголовок, необязательно",
    "footer": "подпись внизу слайдов, по умолчанию название",
    "theme": "sand",
    "slides": [
      {"title": "Название раздела"},
      {"title": "Заголовок", "text": "абзац текста"},
      {"title": "Заголовок", "bullets": ["пункт", "ещё пункт"]},
      {"title": "Заголовок", "bullets": ["пункт"], "image": "/путь/картинка.png"},
      {"title": "Заголовок", "image": "/путь/фото.jpg", "notes": "текст докладчику"}
    ]
  }

Поля слайда: title, subtitle, text, bullets (список строк), image (путь к
файлу), notes (заметки докладчика). Любое можно опустить. Слайд, где есть
только title, становится разделительной плашкой.

Темы: sand (по умолчанию, тёплая светлая), night (тёмная), mint (светлая
зелёная). Своих цветов не выдумывать, брать из этих трёх.

Запуск:
  python3 make_deck.py spec.json [--out файл.pptx] [--theme night] [--json]
  cat spec.json | python3 make_deck.py -

Куда кладётся файл по умолчанию: ~/.hermes/cache/documents. Это одна из папок,
из которых гейтвею разрешено отправлять вложения, поэтому файл оттуда уходит в
Telegram без плясок. Отправить в чат:
  hermes send --to telegram "MEDIA:<путь к файлу>"
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt
except ImportError:  # pragma: no cover
    sys.stderr.write(
        "Нет библиотеки python-pptx. Поставить: pip install python-pptx\n"
    )
    raise SystemExit(3)

# Оформление. Три готовых темы, свои цвета не придумывать.
THEMES = {
    "sand": {"bg": "FBF7F2", "ink": "1F1B16", "accent": "C0703C", "muted": "6E6259"},
    "night": {"bg": "14161A", "ink": "F2F3F5", "accent": "7FB3FF", "muted": "9AA3AE"},
    "mint": {"bg": "F4FAF7", "ink": "14231C", "accent": "2E9E6B", "muted": "5C7268"},
}
# Arial есть и на маке, и в Windows, и в Google Slides, и с кириллицей.
FONT = "Arial"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.9)
BODY_W = SLIDE_W - 2 * MARGIN

TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e",
    "ж": "zh", "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m",
    "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u",
    "ф": "f", "х": "h", "ц": "c", "ч": "ch", "ш": "sh", "щ": "sch", "ъ": "",
    "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
}


def slugify(text, fallback="deck"):
    """Имя файла латиницей: кириллица переводится, пробелы в дефисы."""
    out = []
    for ch in (text or "").lower():
        out.append(TRANSLIT.get(ch, ch))
    slug = re.sub(r"[^a-z0-9]+", "-", "".join(out)).strip("-")
    return (slug or fallback)[:60]


def fail(message):
    sys.stderr.write(message.rstrip() + "\n")
    raise SystemExit(2)


def load_spec(source):
    """Читает описание из файла или со стандартного входа и проверяет форму."""
    try:
        raw = sys.stdin.read() if source == "-" else open(source, encoding="utf-8").read()
    except OSError as exc:
        fail("Не смог прочитать описание слайдов: %s" % exc)
    try:
        spec = json.loads(raw)
    except json.JSONDecodeError as exc:
        fail("Описание слайдов не разобралось как JSON: %s" % exc)
    if not isinstance(spec, dict):
        fail("Описание слайдов должно быть объектом JSON с полем slides.")
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        fail("В описании нет непустого списка slides.")
    for i, item in enumerate(slides, 1):
        if not isinstance(item, dict):
            fail("Слайд %d описан не объектом JSON." % i)
        bullets = item.get("bullets")
        if bullets is not None and not isinstance(bullets, list):
            fail("Слайд %d: bullets должен быть списком строк." % i)
    theme = spec.get("theme") or "sand"
    if theme not in THEMES:
        fail("Тема «%s» неизвестна. Есть: %s" % (theme, ", ".join(sorted(THEMES))))
    return spec


def rgb(value):
    return RGBColor.from_string(value)


def new_slide(prs, theme):
    """Пустой слайд, залитый фоном темы."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(theme["bg"])
    return slide


def textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    return frame


def write(frame, lines, size, color, bold=False, spacing=1.2, space_after=8,
          bullet=False, first=False):
    """Пишет абзацы в готовый текстовый блок. bullet: точка и висячий отступ."""
    for i, line in enumerate(lines):
        par = frame.paragraphs[0] if (first and i == 0) else frame.add_paragraph()
        if bullet:
            # marL и indent это атрибуты, а не вложенные теги: порядок элементов
            # в pPr не ломается, файл остаётся валидным.
            props = par._p.get_or_add_pPr()
            props.set("marL", str(Inches(0.32)))
            props.set("indent", str(-Inches(0.32)))
        run = par.add_run()
        run.text = ("•  " + line) if bullet else line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = FONT
        run.font.color.rgb = rgb(color)
        par.line_spacing = spacing
        par.space_after = Pt(space_after)


def accent_rule(slide, theme, left, top, width=Inches(1.4), height=Inches(0.055)):
    """Тонкая цветная линия под заголовком."""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(theme["accent"])
    bar.line.fill.background()
    bar.shadow.inherit = False
    # У новой фигуры остаётся ссылка на стиль темы с эффектом тени (effectRef).
    # LibreOffice рисует по ней тень даже при пустом effectLst, поэтому блок
    # стиля убираем целиком: цвет и отсутствие рамки заданы явно выше.
    style = bar._element.find(qn("p:style"))
    if style is not None:
        bar._element.remove(style)
    return bar


def place_image(slide, path, left, top, width, height):
    """Вписывает картинку в прямоугольник, пропорции сохраняются."""
    full = os.path.expanduser(str(path))
    if not os.path.isfile(full):
        sys.stderr.write("ВНИМАНИЕ: картинка не найдена, слайд собран без неё: %s\n" % path)
        return False
    try:
        pic = slide.shapes.add_picture(full, left, top, width=width)
    except Exception as exc:  # формат, который python-pptx не читает
        sys.stderr.write("ВНИМАНИЕ: картинку не удалось вставить (%s): %s\n" % (exc, path))
        return False
    if pic.height > height:
        ratio = height / pic.height
        pic.height = int(pic.height * ratio)
        pic.width = int(pic.width * ratio)
    pic.left = int(left + (width - pic.width) / 2)
    pic.top = int(top + (height - pic.height) / 2)
    return True


def cover_slide(prs, spec, theme):
    slide = new_slide(prs, theme)
    accent_rule(slide, theme, MARGIN, Inches(2.35), width=Inches(2.2))
    frame = textbox(slide, MARGIN, Inches(2.75), BODY_W, Inches(2.2))
    write(frame, [spec.get("title") or "Презентация"], 44, theme["ink"], bold=True,
          spacing=1.05, space_after=0, first=True)
    subtitle = (spec.get("subtitle") or "").strip()
    if subtitle:
        sub = textbox(slide, MARGIN, Inches(4.55), BODY_W, Inches(1.0))
        write(sub, [subtitle], 18, theme["muted"], spacing=1.2, space_after=0, first=True)
    stamp = textbox(slide, MARGIN, SLIDE_H - Inches(1.0), BODY_W, Inches(0.4))
    write(stamp, [dt.date.today().strftime("%d.%m.%Y")], 12, theme["muted"],
          spacing=1.0, space_after=0, first=True)
    return slide


def section_slide(prs, item, theme):
    """Слайд-разделитель: есть только заголовок."""
    slide = new_slide(prs, theme)
    accent_rule(slide, theme, MARGIN, Inches(3.0), width=Inches(1.8))
    frame = textbox(slide, MARGIN, Inches(3.4), BODY_W, Inches(1.6))
    write(frame, [item.get("title") or ""], 34, theme["ink"], bold=True,
          spacing=1.1, space_after=0, first=True)
    subtitle = (item.get("subtitle") or "").strip()
    if subtitle:
        sub = textbox(slide, MARGIN, Inches(4.6), BODY_W, Inches(0.9))
        write(sub, [subtitle], 16, theme["muted"], spacing=1.2, space_after=0, first=True)
    return slide


def content_slide(prs, item, theme, number, footer):
    slide = new_slide(prs, theme)
    title = (item.get("title") or "").strip()
    subtitle = (item.get("subtitle") or "").strip()
    text = (item.get("text") or "").strip()
    bullets = [str(b).strip() for b in (item.get("bullets") or []) if str(b).strip()]
    image = (item.get("image") or "").strip()

    top = Inches(0.72)
    if title:
        frame = textbox(slide, MARGIN, top, BODY_W, Inches(1.0))
        write(frame, [title], 28, theme["ink"], bold=True, spacing=1.1,
              space_after=0, first=True)
        top = top + Inches(0.85)
        accent_rule(slide, theme, MARGIN, top)
        top = top + Inches(0.42)
    if subtitle:
        sub = textbox(slide, MARGIN, top, BODY_W, Inches(0.5))
        write(sub, [subtitle], 15, theme["muted"], spacing=1.2, space_after=0, first=True)
        top = top + Inches(0.6)

    body_h = SLIDE_H - top - Inches(0.95)
    has_text = bool(text or bullets)
    if image and has_text:
        gap = Inches(0.5)
        text_w = int((BODY_W - gap) * 0.55)
        place_image(slide, image, MARGIN + text_w + gap, top,
                    BODY_W - text_w - gap, body_h)
    elif image:
        text_w = BODY_W
        place_image(slide, image, MARGIN, top, BODY_W, body_h)
    else:
        text_w = BODY_W

    if has_text:
        frame = textbox(slide, MARGIN, top, text_w, body_h)
        frame.vertical_anchor = MSO_ANCHOR.TOP
        first = True
        if text:
            paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
            write(frame, paragraphs, 18, theme["ink"], spacing=1.25, space_after=12,
                  first=first)
            first = False
        if bullets:
            write(frame, bullets, 18, theme["ink"], spacing=1.2, space_after=10,
                  bullet=True, first=first)

    foot = textbox(slide, MARGIN, SLIDE_H - Inches(0.72), BODY_W - Inches(0.6), Inches(0.35))
    write(foot, [footer], 11, theme["muted"], spacing=1.0, space_after=0, first=True)
    num = textbox(slide, SLIDE_W - MARGIN - Inches(0.6), SLIDE_H - Inches(0.72),
                  Inches(0.6), Inches(0.35))
    write(num, [str(number)], 11, theme["muted"], spacing=1.0, space_after=0, first=True)
    num.paragraphs[0].alignment = PP_ALIGN.RIGHT

    notes = (item.get("notes") or "").strip()
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def default_out_dir():
    """Папка, из которой гейтвею разрешено отправлять вложения."""
    home = os.environ.get("HERMES_HOME") or os.path.join(os.path.expanduser("~"), ".hermes")
    return os.path.join(home, "cache", "documents")


def build(spec, out_path=None, theme_name=None):
    theme = THEMES[theme_name or spec.get("theme") or "sand"]
    footer = (spec.get("footer") or spec.get("title") or "").strip()

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    if (spec.get("title") or "").strip():
        cover_slide(prs, spec, theme)

    number = 1
    for item in spec["slides"]:
        bare = not any(str(item.get(k) or "").strip() for k in ("text", "image", "subtitle")) \
            and not [b for b in (item.get("bullets") or []) if str(b).strip()]
        if bare and (item.get("title") or "").strip():
            section_slide(prs, item, theme)
        else:
            content_slide(prs, item, theme, number, footer)
        number += 1

    if not out_path:
        stamp = dt.date.today().strftime("%Y-%m-%d")
        name = "%s-%s.pptx" % (slugify(spec.get("title"), "prezentaciya"), stamp)
        out_path = os.path.join(default_out_dir(), name)
    out_path = os.path.expanduser(out_path)
    directory = os.path.dirname(os.path.abspath(out_path))
    os.makedirs(directory, exist_ok=True)
    prs.save(out_path)
    return os.path.abspath(out_path), len(prs.slides)


def main():
    ap = argparse.ArgumentParser(description="Презентация .pptx из описания слайдов")
    ap.add_argument("spec", help="файл с описанием слайдов или - для стандартного входа")
    ap.add_argument("--out", help="куда сохранить .pptx")
    ap.add_argument("--theme", choices=sorted(THEMES), help="тема оформления")
    ap.add_argument("--json", action="store_true", help="ответ машинным JSON")
    args = ap.parse_args()

    spec = load_spec(args.spec)
    path, slides = build(spec, args.out, args.theme)
    if args.json:
        print(json.dumps({"path": path, "slides": slides}, ensure_ascii=False))
    else:
        print(path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
